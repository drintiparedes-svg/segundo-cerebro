"""Lente financiera por proyecto: presupuesto, ejecución, unidades
económicas e impacto. Puro cálculo local sobre tus planillas.

- Presupuesto: hoja «Presupuesto» del plan Excel (Fase · Ítem · Detalle ·
  Costo unit. · Cantidad · Total, con subtotal/contingencia/total).
- Ejecución: CSV/XLSX de gasto real (fecha, ítem, fase, monto).
- Unidades económicas: costo total por unidad definida en projects.md
  (p. ej. «kit de autotoma», «paciente tamizado») y por fase.
- Impacto: beneficio − costo, ROI, payback y sensibilidad ±20 %.
Todo queda en .brain/finance/<proyecto>/.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import asdict
from datetime import date
from pathlib import Path

NUM_RE = re.compile(r"[-+]?\d[\d.,]*")
TOTAL_WORDS = ("total", "subtotal", "contingencia", "contingency")


def _num(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = NUM_RE.search(str(value).replace("$", ""))
    if not m:
        return None
    txt = m.group(0)
    # el último separador es decimal si no le siguen exactamente 3 dígitos
    last = max(txt.rfind(","), txt.rfind("."))
    if last >= 0 and len(txt) - last - 1 != 3:
        txt = txt[:last].replace(",", "").replace(".", "") + "." + txt[last + 1:]
    else:
        txt = txt.replace(",", "").replace(".", "")
    try:
        return float(txt)
    except ValueError:
        return None


def finance_dir(brain_dir: str | Path, project_id: str) -> Path:
    d = Path(brain_dir) / "finance" / project_id
    d.mkdir(parents=True, exist_ok=True)
    return d


# ── presupuesto ───────────────────────────────────────────────────────────

HEADERS = {
    "phase": ("fase", "phase", "etapa"),
    "item": ("ítem", "item", "concepto", "partida"),
    "detail": ("unidad", "detalle", "detail", "descripci"),
    "unit_cost": ("costo unit", "precio unit", "unit cost", "valor unit"),
    "qty": ("cantidad", "qty", "quantity", "n°", "nº"),
    "total": ("total", "monto", "amount"),
}


def _match_header(cells: list) -> dict | None:
    low = [str(c).strip().lower() if c is not None else "" for c in cells]
    if not any(any(n in c for n in HEADERS["item"]) for c in low):
        return None
    cols = {}
    for key, names in HEADERS.items():
        for i, c in enumerate(low):
            if c and any(c.startswith(n) or n in c for n in names) and key not in cols:
                cols[key] = i
                break
    return cols if "item" in cols and ("total" in cols or "unit_cost" in cols) else None


def parse_budget_xlsx(path: str | Path, sheet_hint: str = "presupuesto") -> dict:
    """Devuelve {items, subtotal, contingency, total, currency, sheet}."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    best = None
    try:
        sheets = sorted(wb.worksheets, key=lambda ws: 0 if sheet_hint in ws.title.lower() else 1)
        for ws in sheets:
            cols, items, totals = None, [], {}
            for raw in ws.iter_rows(values_only=True):
                cells = list(raw)
                if cols is None:
                    cols = _match_header(cells)
                    continue
                get = lambda k: (cells[cols[k]] if k in cols and cols[k] < len(cells) else None)
                label = " ".join(str(c) for c in cells if c is not None).lower()
                total = _num(get("total"))
                if any(w in label for w in TOTAL_WORDS) and not get("qty"):
                    key = "contingency" if "conting" in label else ("subtotal" if "subtotal" in label else "total")
                    if total is not None:
                        totals[key] = total
                    continue
                item = get("item")
                if item is None or not str(item).strip():
                    continue
                unit_cost, qty = _num(get("unit_cost")), _num(get("qty"))
                if total is None and unit_cost is not None and qty is not None:
                    total = unit_cost * qty
                if total is None:
                    continue
                items.append({"phase": str(get("phase") or "").strip() or None, "item": str(item).strip(),
                              "detail": str(get("detail") or "").strip() or None,
                              "unit_cost": unit_cost, "qty": qty, "total": round(total, 2)})
            if items and (best is None or len(items) > len(best["items"])):
                subtotal = totals.get("subtotal", round(sum(i["total"] for i in items), 2))
                best = {"items": items, "subtotal": subtotal,
                        "contingency": totals.get("contingency", 0.0),
                        "total": totals.get("total", subtotal + totals.get("contingency", 0.0)),
                        "currency": "CLP", "sheet": ws.title, "path": str(path)}
    finally:
        wb.close()
    return best or {"items": [], "subtotal": 0.0, "contingency": 0.0, "total": 0.0, "currency": "CLP", "sheet": None, "path": str(path)}


def import_budget(brain_dir: str | Path, project_id: str, path: str | Path) -> dict:
    budget = parse_budget_xlsx(path)
    budget["imported_at"] = date.today().isoformat()
    (finance_dir(brain_dir, project_id) / "budget.json").write_text(
        json.dumps(budget, ensure_ascii=False, indent=2), encoding="utf-8")
    return budget


def import_actuals(brain_dir: str | Path, project_id: str, path: str | Path) -> dict:
    """CSV/XLSX con columnas fecha, ítem/concepto, fase (opcional), monto."""
    p = Path(path)
    rows = []
    if p.suffix.lower() == ".csv":
        with p.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f, delimiter=";" if ";" in f.readline() else ",")
            f.seek(0)
            reader = csv.DictReader(f, delimiter=reader.reader.dialect.delimiter if hasattr(reader.reader, "dialect") else ",")
            raw_rows = list(reader)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(str(p), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(it)]
        raw_rows = [dict(zip(header, r)) for r in it]
        wb.close()
    for r in raw_rows:
        low = {k.strip().lower(): v for k, v in r.items() if k}
        amount = next((_num(v) for k, v in low.items() if any(n in k for n in ("monto", "amount", "total", "gasto"))), None)
        if amount is None:
            continue
        rows.append({
            "date": next((str(v)[:10] for k, v in low.items() if "fecha" in k or "date" in k), None),
            "item": next((str(v) for k, v in low.items() if any(n in k for n in ("ítem", "item", "concepto", "detalle"))), ""),
            "phase": next((str(v) for k, v in low.items() if "fase" in k or "phase" in k), None),
            "amount": amount,
        })
    data = {"rows": rows, "total": round(sum(r["amount"] for r in rows), 2),
            "imported_at": date.today().isoformat(), "path": str(p)}
    (finance_dir(brain_dir, project_id) / "actuals.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data


def load_json(brain_dir: str | Path, project_id: str, name: str) -> dict | None:
    p = finance_dir(brain_dir, project_id) / f"{name}.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


# ── unidades económicas e impacto ─────────────────────────────────────────

def by_phase(budget: dict, actuals: dict | None) -> dict:
    out: dict[str, dict] = {}
    for i in budget.get("items", []):
        ph = i.get("phase") or "—"
        out.setdefault(ph, {"budget": 0.0, "spent": 0.0})["budget"] += i["total"]
    for r in (actuals or {}).get("rows", []):
        ph = r.get("phase") or "—"
        out.setdefault(ph, {"budget": 0.0, "spent": 0.0})["spent"] += r["amount"]
    for ph, v in out.items():
        v["budget"] = round(v["budget"], 2)
        v["spent"] = round(v["spent"], 2)
        v["execution_pct"] = round(100 * v["spent"] / v["budget"]) if v["budget"] else None
    return out


def unit_economics(project, budget: dict, actuals: dict | None = None, indicators: dict | None = None) -> dict:
    total = float(budget.get("total") or 0.0)
    spent = float((actuals or {}).get("total") or 0.0)
    units = []
    for u in getattr(project, "units", []) or []:
        vol = float(u.get("volume") or 0)
        if vol <= 0:
            continue
        planned = round(total / vol, 2)
        real = round(spent / vol, 2) if spent else None
        deviation = round(100 * (real - planned) / planned) if (real and planned) else None
        row = {"unit": u.get("name"), "volume": vol, "cost_planned": planned, "cost_real": real,
               "deviation_pct": deviation}
        if u.get("price"):
            price = float(u["price"])
            row["price"] = price
            fixed = float(u.get("fixed_cost") or 0)
            variable = float(u.get("variable_cost") or planned)
            row["breakeven_units"] = round(fixed / (price - variable)) if price > variable and fixed else None
            row["margin_pct"] = round(100 * (price - variable) / price) if price else None
        units.append(row)
    out = {"project": project.id, "currency": budget.get("currency", "CLP"), "budget_total": total,
           "spent": spent, "execution_pct": round(100 * spent / total) if total else None,
           "remaining": round(total - spent, 2), "by_phase": by_phase(budget, actuals), "units": units}
    if indicators:
        from .connectors.indicators import clp_to_uf, clp_to_usd
        out["budget_usd"] = clp_to_usd(total, indicators)
        out["budget_uf"] = clp_to_uf(total, indicators)
    return out


def impact(project, budget: dict, horizon_months: int = 12) -> dict:
    """Beneficio − costo, ROI, payback y sensibilidad ±20 % en volumen y costo."""
    cost = float(budget.get("total") or 0.0)
    benefits = getattr(project, "benefits", []) or []
    annual_benefit = 0.0
    for b in benefits:
        v = float(b.get("value_clp") or 0)
        per = (b.get("per") or "year").lower()
        annual_benefit += v * 12 if per.startswith("month") else v
    if not cost:
        return {"cost": 0.0, "annual_benefit": annual_benefit, "net": None, "roi_pct": None, "payback_months": None, "sensitivity": []}
    net = annual_benefit * horizon_months / 12 - cost
    roi = round(100 * net / cost) if cost else None
    payback = round(cost / (annual_benefit / 12), 1) if annual_benefit > 0 else None
    sens = []
    for label, vol_f, cost_f in (("volumen −20 %", 0.8, 1.0), ("volumen +20 %", 1.2, 1.0),
                                 ("costo +20 %", 1.0, 1.2), ("costo −20 %", 1.0, 0.8)):
        c = cost * cost_f
        b = annual_benefit * vol_f * horizon_months / 12
        sens.append({"scenario": label, "cost": round(c), "benefit": round(b), "net": round(b - c),
                     "roi_pct": round(100 * (b - c) / c) if c else None})
    return {"cost": cost, "annual_benefit": annual_benefit, "horizon_months": horizon_months,
            "net": round(net), "roi_pct": roi, "payback_months": payback, "sensitivity": sens,
            "benefits": benefits}


def alerts(project, ue: dict, opportunities: list | None = None, today: date | None = None) -> list[dict]:
    today = today or date.today()
    out = []
    for ph, v in ue.get("by_phase", {}).items():
        if v.get("execution_pct") is not None and v["execution_pct"] >= 90:
            out.append({"level": "L1", "kind": "ejecución", "text": f"Fase {ph}: {v['execution_pct']} % del presupuesto ejecutado"})
    for u in ue.get("units", []):
        if u.get("deviation_pct") is not None and abs(u["deviation_pct"]) >= 15:
            out.append({"level": "L1", "kind": "costo unitario",
                        "text": f"{u['unit']}: costo real {u['cost_real']:,.0f} vs plan {u['cost_planned']:,.0f} ({u['deviation_pct']:+d} %)"})
    for ko in opportunities or []:
        if ko.valid_to and 0 <= (date.fromisoformat(ko.valid_to) - today).days <= 14:
            out.append({"level": "L1", "kind": "financiamiento", "text": f"Cierra {ko.valid_to}: {ko.title}"})
    return out


def report(project, ue: dict, imp: dict, alerts_list: list[dict], opportunities: list | None = None) -> str:
    cur = ue.get("currency", "CLP")
    fmt = lambda v: f"{v:,.0f} {cur}" if isinstance(v, (int, float)) else "—"
    lines = [f"# Lente financiera · {project.name}", ""]
    lines.append(f"- Presupuesto: **{fmt(ue['budget_total'])}**"
                 + (f" (≈ {ue['budget_usd']:,.0f} USD · {ue['budget_uf']:,.1f} UF)" if ue.get("budget_usd") else ""))
    lines.append(f"- Ejecutado: {fmt(ue['spent'])}" + (f" ({ue['execution_pct']} %)" if ue.get("execution_pct") is not None else "")
                 + f" · restante {fmt(ue['remaining'])}")
    if ue.get("by_phase"):
        lines += ["", "## Por fase", "| Fase | Presupuesto | Ejecutado | % |", "|---|---:|---:|---:|"]
        lines += [f"| {ph} | {fmt(v['budget'])} | {fmt(v['spent'])} | {v['execution_pct'] if v['execution_pct'] is not None else '—'} |"
                  for ph, v in ue["by_phase"].items()]
    if ue.get("units"):
        lines += ["", "## Unidades económicas", "| Unidad | Volumen | Costo plan | Costo real | Desvío |", "|---|---:|---:|---:|---:|"]
        for u in ue["units"]:
            lines.append(f"| {u['unit']} | {u['volume']:,.0f} | {fmt(u['cost_planned'])} | {fmt(u['cost_real']) if u['cost_real'] else '—'} | "
                         f"{str(u['deviation_pct']) + ' %' if u['deviation_pct'] is not None else '—'} |")
            if u.get("breakeven_units"):
                lines.append(f"  - Punto de equilibrio: {u['breakeven_units']:,} unidades (margen {u['margin_pct']} %)")
    lines += ["", "## Impacto"]
    if imp.get("net") is None:
        lines.append("- Sin presupuesto o beneficios definidos (edita `benefits` en projects.md).")
    else:
        lines.append(f"- Beneficio anual estimado: {fmt(imp['annual_benefit'])} · horizonte {imp['horizon_months']} meses")
        lines.append(f"- Neto: **{fmt(imp['net'])}** · ROI {imp['roi_pct']} %"
                     + (f" · payback {imp['payback_months']} meses" if imp.get("payback_months") else ""))
        lines += ["", "| Escenario | Costo | Beneficio | Neto | ROI |", "|---|---:|---:|---:|---:|"]
        lines += [f"| {s['scenario']} | {fmt(s['cost'])} | {fmt(s['benefit'])} | {fmt(s['net'])} | {s['roi_pct']} % |" for s in imp["sensitivity"]]
        for b in imp.get("benefits", []):
            lines.append(f"  - {b.get('name')}: {fmt(float(b.get('value_clp') or 0))}/{b.get('per', 'year')} — {b.get('basis', 'supuesto')}")
    if opportunities:
        lines += ["", "## Financiamiento abierto que calza"]
        lines += [f"- cierra {o.valid_to} · {o.title}" for o in opportunities]
    if alerts_list:
        lines += ["", "## Alertas"]
        lines += [f"- ⚠ [{a['kind']}] {a['text']}" for a in alerts_list]
    lines += ["", "---", "_Cálculo local sobre tus planillas; supuestos editables en brain/self/projects.md._"]
    return "\n".join(lines)


def project_finance(store, brain_dir: str | Path, project, indicators: dict | None = None,
                    today: date | None = None) -> dict:
    budget = load_json(brain_dir, project.id, "budget") or {"items": [], "total": 0.0, "currency": "CLP"}
    actuals = load_json(brain_dir, project.id, "actuals")
    ue = unit_economics(project, budget, actuals, indicators)
    imp = impact(project, budget)
    opps = [k for k in store.list_knowledge_objects(ko_type="opportunity", status="active", limit=200)
            if k.project == project.name or f"proyecto:{project.id}" in (k.tags or [])]
    al = alerts(project, ue, opps, today)
    return {"project": {"id": project.id, "name": project.name}, "unit_economics": ue, "impact": imp,
            "alerts": al, "opportunities": [asdict(o) for o in opps], "has_budget": bool(budget.get("items")),
            "markdown": report(project, ue, imp, al, opps)}
