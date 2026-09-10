"""Proyectos especiales (la tesis, un producto, una postulación): hitos,
atrasos y próximos pasos como parte de la gestión activa.

- brain/self/projects.md: frontmatter editable a mano con cada proyecto
  (área, deadline, hitos, personas).
- `import_plan`: convierte un plan de trabajo en Excel (actividades por
  semana o con fechas) en compromisos con fecha de vencimiento, ligados al
  archivo original como fuente. Reimportar reemplaza, no duplica.
- `project_brief` / `week_review`: secciones para `sb today` y `sb week`.
Todo local, sin llamadas externas.
"""

from __future__ import annotations

import hashlib
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import yaml

from .areas import FRONTMATTER_RE
from .models import KnowledgeObject

DEFAULT_PROJECTS_FILE = os.environ.get("SB_PROJECTS", "brain/self/projects.md")
DONE_WORDS = ("complet", "hecho", "listo", "cerrad", "done", "finaliz", "termin")
WEEK_RE = re.compile(r"[Ss](?:em\.?)?\s*(\d{1,2})")


@dataclass
class Milestone:
    name: str
    due: str            # YYYY-MM-DD


@dataclass
class Project:
    id: str
    name: str
    area: str | None = None
    deadline: str | None = None
    start: str | None = None
    milestones: list[Milestone] = field(default_factory=list)
    people: list[str] = field(default_factory=list)
    keywords: list[str] = field(default_factory=list)


def _iso(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def load_projects(path: str | Path = DEFAULT_PROJECTS_FILE) -> list[Project]:
    p = Path(path)
    if not p.exists():
        return []
    m = FRONTMATTER_RE.match(p.read_text(encoding="utf-8"))
    if not m:
        return []
    try:
        data = yaml.safe_load(m.group(1)) or {}
    except yaml.YAMLError:
        return []
    projects = []
    for raw in data.get("projects", []):
        if not raw.get("id") or not raw.get("name"):
            continue
        milestones = [Milestone(name=str(ms.get("name")), due=_iso(ms.get("due")) or "")
                      for ms in raw.get("milestones", []) or [] if ms.get("name")]
        projects.append(Project(
            id=str(raw["id"]), name=str(raw["name"]), area=raw.get("area"),
            deadline=_iso(raw.get("deadline")), start=_iso(raw.get("start")),
            milestones=[ms for ms in milestones if ms.due],
            people=[str(x) for x in raw.get("people", []) or []],
            keywords=[str(k).lower() for k in raw.get("keywords", []) or []],
        ))
    return projects


# ── plan de trabajo en Excel → compromisos con fecha ─────────────────────

HEADER_KEYS = {
    "activity": ("actividad", "tarea", "activity", "task"),
    "week": ("sem", "semana", "week"),
    "start": ("inicio", "start"),
    "end": ("fin", "término", "termino", "vencimiento", "fecha", "due", "end"),
    "status": ("estado", "status"),
    "responsible": ("responsable", "owner"),
    "deliverable": ("entregable", "deliverable"),
    "priority": ("prioridad", "priority"),
    "phase": ("fase", "phase"),
}


def _match_header(cells: list) -> dict | None:
    low = [str(c).strip().lower() if c is not None else "" for c in cells]
    if not any("activ" in c or "tarea" in c or "task" in c for c in low):
        return None
    cols = {}
    for key, names in HEADER_KEYS.items():
        for i, c in enumerate(low):
            if c and any(c.startswith(n) or n in c for n in names) and key not in cols:
                cols[key] = i
                break
    return cols if "activity" in cols and len(cols) >= 2 else None


def parse_plan_xlsx(path: str | Path, start: str | None = None) -> list[dict]:
    """Recorre todas las hojas y se queda con la que más actividades con
    fecha/semana aporta (p. ej. «Actividades Detalladas» antes que un Gantt
    o una portada). Las semanas (S1, S2…) se convierten a fecha si se
    conoce `start`."""
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    start_date = date.fromisoformat(start) if start else None
    best: list[dict] = []
    try:
        for ws in wb.worksheets:
            rows: list[dict] = []
            cols = None
            for raw in ws.iter_rows(values_only=True):
                cells = list(raw)
                if cols is None:
                    cols = _match_header(cells)
                    continue
                get = lambda key: (cells[cols[key]] if key in cols and cols[key] < len(cells) else None)
                activity = get("activity")
                if activity is None or not str(activity).strip():
                    continue
                activity = str(activity).strip()
                if activity.isupper() and len(activity.split()) > 2:
                    continue  # títulos de fase en mayúsculas
                due = _iso(get("end")) or None
                week = None
                wk = get("week")
                if wk is not None:
                    m = WEEK_RE.search(str(wk)) or re.search(r"\d{1,2}", str(wk))
                    if m:
                        week = int(m.group(1) if m.lastindex else m.group(0))
                if due is None and week and start_date:
                    due = (start_date + timedelta(days=week * 7 - 1)).isoformat()
                status_raw = str(get("status") or "").lower()
                rows.append({
                    "activity": activity, "phase": str(get("phase") or "").strip() or None,
                    "week": week, "start": _iso(get("start")), "due": due,
                    "status": "done" if any(w in status_raw for w in DONE_WORDS) else "active",
                    "status_raw": status_raw or None,
                    "responsible": str(get("responsible") or "").strip() or None,
                    "deliverable": str(get("deliverable") or "").strip() or None,
                    "priority": str(get("priority") or "").strip() or None,
                    "sheet": ws.title,
                })
            key = lambda rs: (sum(1 for r in rs if r["due"] or r["week"]), len(rs))
            if rows and key(rows) > key(best):
                best = rows
    finally:
        wb.close()
    return best


def import_plan(store, path: str | Path, project: Project, start: str | None = None,
                area: str | None = None) -> dict:
    """Ingesta el Excel como documento (si no estaba) y crea/reemplaza los
    compromisos del plan. Ids deterministas: reimportar no duplica."""
    from .connectors.localfs import file_to_document, read_file_text

    path = Path(path).expanduser()
    start = start or project.start
    rows = parse_plan_xlsx(path, start)
    if not rows:
        return {"error": "no encontré una hoja con columna «Actividad»", "tasks": 0}

    text = read_file_text(path) or path.name
    doc = file_to_document(path, "plan-de-trabajo", text)
    doc.doc_type = "plan"
    doc.area = area or project.area
    doc.metadata["project"] = project.name
    if not store.add_document(doc):
        existing = next((d for d in store.list_documents(limit=100_000)
                         if d.path == doc.path), None)
        doc = existing or doc
        store.delete_derived(doc.id)
    store.mark_extractor(doc.id, "PlanImporter")

    created = 0
    for row in rows:
        digest = hashlib.sha1(f"{project.id}|{row['activity']}|{row['week']}".encode()).hexdigest()[:12]
        tags = [t for t in (row["phase"], row["priority"] and f"prioridad:{row['priority'].lower()}") if t]
        detail = f" — entregable: {row['deliverable']}" if row["deliverable"] else ""
        store.add_knowledge_object(KnowledgeObject(
            id=f"ko-plan-{digest}", ko_type="task", title=row["activity"][:80],
            statement=f"{row['activity']}{detail}",
            date=row["start"] or row["due"] or doc.date,
            people=[p for p in project.people if row["responsible"] in (None, "", "IP")] or [],
            project=project.name, status=row["status"], confidence="confirmed",
            source_doc=doc.id, tags=tags, valid_from=row["start"],
            valid_to=row["due"], area=area or project.area,
        ))
        created += 1
    dated = sum(1 for r in rows if r["due"])
    return {"tasks": created, "dated": dated, "doc_id": doc.id, "path": str(path),
            "warning": None if dated or not start is None else
            "sin fechas: pasa --start YYYY-MM-DD para convertir las semanas en vencimientos"}


# ── señales para el brief y la revisión semanal ───────────────────────────

def project_tasks(store, project: Project) -> list[KnowledgeObject]:
    return store.list_knowledge_objects(ko_type="task", project=project.name, limit=1000)


def project_status(store, project: Project, today: date | None = None) -> dict:
    today = today or date.today()
    week = str(today + timedelta(days=7))
    tasks = project_tasks(store, project)
    active = [t for t in tasks if t.status == "active"]
    overdue = sorted((t for t in active if t.valid_to and t.valid_to < str(today)),
                     key=lambda t: t.valid_to)
    due_soon = sorted((t for t in active if t.valid_to and str(today) <= t.valid_to <= week),
                      key=lambda t: t.valid_to)
    upcoming = sorted((m for m in project.milestones if m.due >= str(today)),
                      key=lambda m: m.due)
    questions = store.list_knowledge_objects(ko_type="question", status="active",
                                             project=project.name, limit=20)
    done = sum(1 for t in tasks if t.status == "done")
    return {"project": project, "tasks": len(tasks), "done": done, "active": len(active),
            "overdue": overdue, "due_soon": due_soon,
            "next_milestone": upcoming[0] if upcoming else None,
            "days_to_deadline": ((date.fromisoformat(project.deadline) - today).days
                                 if project.deadline else None),
            "questions": questions}


def project_brief(store, projects: list[Project], today: date | None = None) -> list[str]:
    """Líneas Markdown de la sección «Proyectos especiales»."""
    today = today or date.today()
    lines = []
    for project in projects:
        st = project_status(store, project, today)
        head = f"### {project.name}"
        bits = []
        if st["days_to_deadline"] is not None:
            bits.append(f"entrega en {st['days_to_deadline']} días ({project.deadline})")
        if st["tasks"]:
            bits.append(f"{st['done']}/{st['tasks']} actividades hechas")
        lines.append(head + (" — " + " · ".join(bits) if bits else ""))
        if st["next_milestone"]:
            ms = st["next_milestone"]
            days = (date.fromisoformat(ms.due) - today).days
            lines.append(f"- Próximo hito: **{ms.name}** en {days} días ({ms.due})")
        for t in st["overdue"][:5]:
            lines.append(f"- ⚠ Atrasada desde {t.valid_to}: {t.title}")
        for t in st["due_soon"][:5]:
            lines.append(f"- Vence {t.valid_to}: {t.title}")
        for q in st["questions"][:2]:
            lines.append(f"- [question] {q.statement}")
        if not (st["next_milestone"] or st["overdue"] or st["due_soon"] or st["questions"]):
            lines.append("- Sin hitos ni vencimientos próximos registrados.")
        lines.append("")
    return lines


def week_review(store, brain_dir: str | Path, projects: list[Project],
                area_names: dict | None = None, today: date | None = None) -> str:
    """Revisión semanal: qué se decidió, qué se cerró, qué se atrasó y qué
    viene. 100% local."""
    import json
    today = today or date.today()
    area_names = area_names or {}
    week_ago = str(today - timedelta(days=7))
    next_week = str(today + timedelta(days=7))
    two_weeks = str(today + timedelta(days=14))
    month_ago = str(today - timedelta(days=30))
    name_of = lambda aid: area_names.get(aid, "Sin área")

    lines = [f"# Revisión semanal — {today.isoformat()}", ""]

    decisions = [d for d in store.list_knowledge_objects(ko_type="decision", limit=500)
                 if d.date >= week_ago]
    lines.append("## Decisiones de la semana")
    lines += [f"- {d.statement} [{name_of(d.area)}]" for d in decisions] or ["- Ninguna registrada."]
    lines.append("")

    tasks = store.list_knowledge_objects(ko_type="task", limit=2000)
    closed = [t for t in tasks if t.status == "done" and t.date >= week_ago]
    overdue = sorted((t for t in tasks if t.status == "active" and t.valid_to
                      and t.valid_to < str(today)), key=lambda t: t.valid_to)
    due = sorted((t for t in tasks if t.status == "active" and t.valid_to
                  and str(today) <= t.valid_to <= next_week), key=lambda t: t.valid_to)
    lines.append("## Compromisos")
    lines.append(f"- Cerrados esta semana: {len(closed)}")
    lines.append(f"- Atrasados: {len(overdue)}")
    lines += [f"  - ⚠ {t.valid_to} · {t.title} [{name_of(t.area)}]" for t in overdue[:10]]
    lines.append(f"- Vencen la próxima semana: {len(due)}")
    lines += [f"  - {t.valid_to} · {t.title} [{name_of(t.area)}]" for t in due[:10]]
    lines.append("")

    lines.append("## Hitos próximos (14 días)")
    hits = []
    for p in projects:
        for ms in p.milestones:
            if str(today) <= ms.due <= two_weeks:
                hits.append(f"- {ms.due} · **{ms.name}** ({p.name})")
    lines += hits or ["- Ninguno."]
    lines.append("")

    stale = [q for q in store.list_knowledge_objects(ko_type="question", status="active", limit=200)
             if q.date <= month_ago]
    if stale:
        lines.append("## Preguntas sin resolver hace más de 30 días")
        lines += [f"- {q.statement} ({q.date}) [{name_of(q.area)}]" for q in stale[:10]]
        lines.append("")

    triage_path = Path(brain_dir) / "reports" / "latest-triage.json"
    if triage_path.exists():
        mail = json.loads(triage_path.read_text(encoding="utf-8"))
        urgent = [m for m in mail if m.get("priority", 5) <= 2]
        lines.append(f"## Correo: {len(urgent)} P1-P2 en el último triaje")
        lines += [f"- P{m['priority']} · {m.get('subject', '')} — {m.get('from', '')}"
                  for m in urgent[:6]]
        lines.append("")

    lines.append("---")
    lines.append("_Generado localmente desde tu memoria; ninguna llamada externa._")
    return "\n".join(lines)
